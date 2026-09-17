#!/usr/bin/env python3
"""Regenerate data/sessions.js from Poker Tracker.xlsx. Run after editing the spreadsheet, then commit."""
import json, datetime, pathlib
from openpyxl import load_workbook
root = pathlib.Path(__file__).parent
ws = load_workbook(root / "Poker Tracker.xlsx", data_only=True)["Sessions"]
rows = []
for r in ws.iter_rows(min_row=2, values_only=True):
    date, stakes, game, buyin, cash, _profit, hours, _rate, _run, notes, _k = r[:11]
    if buyin is None and cash is None: continue
    rows.append({"date": date.strftime("%Y-%m-%d") if isinstance(date, datetime.datetime) else str(date),
                 "stakes": stakes or "", "game": game or "", "buyin": buyin or 0, "cashout": cash or 0,
                 "hours": hours, "notes": notes or ""})
(root / "data" / "sessions.js").write_text("window.SESSIONS = " + json.dumps(rows, indent=1) + ";\n")
print(f"wrote {len(rows)} sessions")
